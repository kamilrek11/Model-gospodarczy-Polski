"""POST /export?format=csv|xlsx: provenance metadata + parameters + per-year results of a /simulate response."""

import csv
import io
from datetime import UTC, datetime
from typing import Any, Literal

from fastapi import APIRouter, HTTPException, Query, Request, Response
from openpyxl import Workbook
from pydantic import BaseModel

from app.model.params import validate
from app.model.simulate import BANDS

HEADER = ["year", "tax", *BANDS, "baseline"]
PARAM_HEADER = ["parametr", "wartość"]
XLSX_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
_FORMULA_PREFIX = ("=", "+", "-", "@", "\t", "\r")


MAX_BODY_BYTES = 1_000_000
MAX_YEARS = 50
MAX_TAXES = 20
MAX_ASSUMPTIONS = 100
MAX_META_TEXT = 64

META_TITLE = "# Metadane"
SYNTHETIC_NOTE = "DANE SYNTETYCZNE – nie są to dane GUS/MF"
MAX_SOURCE_TEXT = 500
ASSUMPTION_SENTENCE = (
    "Wyniki zależą od niezweryfikowanych stałych o statusie „założenie” (brak źródła), wymienionych powyżej; "
    "należy je traktować jako ilustracyjne."
)
STATUS_LABEL = {"from_data": "z danych", "sourced": "ze źródła", "assumption": "założenie"}


class ExportRequest(BaseModel):
    params: dict[str, Any]
    result: dict[str, Any]
    name: str | None = None


def safe(value: Any) -> Any:
    """Neutralise spreadsheet formula injection in user-supplied text."""
    if isinstance(value, str) and value.startswith(_FORMULA_PREFIX):
        return "'" + value
    return value


def result_rows(result: dict) -> list[list]:
    years_raw, taxes_raw = result.get("years"), result.get("taxes")
    if (isinstance(years_raw, list) and len(years_raw) > MAX_YEARS) or (isinstance(taxes_raw, dict) and len(taxes_raw) > MAX_TAXES):
        raise HTTPException(status_code=422, detail=f"Wynik zbyt duży do eksportu (maks. {MAX_YEARS} lat i {MAX_TAXES} podatków).")
    try:
        years = [int(y) for y in result["years"]]
        taxes, baseline = result["taxes"], result["baseline"]
        names = [*taxes, *(["total"] if "total" in result else [])]
        rows = []
        for tax in names:
            bands = taxes[tax] if tax in taxes else result["total"]
            base = (baseline.get(tax) or {}).get("median")
            for i, year in enumerate(years):
                rows.append([year, tax, *(float(bands[b][i]) for b in BANDS), float(base[i]) if base else None])
        return rows
    except (KeyError, IndexError, TypeError, ValueError, AttributeError) as e:
        raise HTTPException(status_code=422, detail=f"Nieprawidłowy wynik symulacji: {e!r}") from e


def _meta_text(result: dict, key: str) -> str:
    value = result.get(key)
    if value is None:
        return "brak"
    if not isinstance(value, str) or len(value) > MAX_META_TEXT:
        raise HTTPException(status_code=422, detail=f"Nieprawidłowe pole wyniku: {key}")
    return value


def meta_rows(result: dict) -> list[list]:
    """Provenance taken from the /simulate response the frontend sends back; missing fields are tolerated."""
    synthetic = result.get("synthetic")
    if synthetic is not None and not isinstance(synthetic, bool):
        raise HTTPException(status_code=422, detail="Nieprawidłowe pole wyniku: synthetic")
    assumptions = result.get("assumptions")
    names: list[str] = []
    constants: list[list] = []
    if assumptions is not None:
        if not isinstance(assumptions, list) or len(assumptions) > MAX_ASSUMPTIONS:
            raise HTTPException(status_code=422, detail="Nieprawidłowe pole wyniku: assumptions")
        for a in assumptions:
            if not isinstance(a, dict) or not isinstance(a.get("key"), str) or len(a["key"]) > MAX_META_TEXT:
                raise HTTPException(status_code=422, detail="Nieprawidłowe pole wyniku: assumptions")
            status = a.get("status")
            source, year, url = a.get("source") or "", a.get("year"), a.get("url") or ""
            if not all(isinstance(t, str) and len(t) <= MAX_SOURCE_TEXT for t in (source, url)) or not (year is None or (isinstance(year, int) and not isinstance(year, bool))):
                raise HTTPException(status_code=422, detail="Nieprawidłowe pole wyniku: assumptions")
            if status == "assumption":
                names.append(a["key"])
            if status in STATUS_LABEL:
                parts = [STATUS_LABEL[status], source, "" if year is None else str(year), url]
                while parts and parts[-1] == "":
                    parts.pop()
                constants.append([f"stala_{a['key']}", "; ".join(parts)])
    rows = [
        ["dane_z", _meta_text(result, "data_updated_at")],
        ["wersja_modelu", _meta_text(result, "model_version")],
        ["dane_syntetyczne", "nieznane" if synthetic is None else ("TAK" if synthetic else "NIE")],
    ]
    if synthetic:
        rows.append(["uwaga", SYNTHETIC_NOTE])
    rows.append(["wygenerowano", datetime.now(UTC).isoformat(timespec="seconds")])
    rows.append(["zalozenia_niezweryfikowane", "brak informacji" if assumptions is None else (", ".join(names) if names else "brak")])
    rows += constants
    if names:
        rows.append(["zastrzezenie", f"{ASSUMPTION_SENTENCE} Założenia: {', '.join(names)}."])
    return rows


def param_rows(req: ExportRequest) -> list[list]:
    try:
        full = validate(req.params)
    except ValueError as e:
        raise HTTPException(status_code=422, detail=str(e)) from e
    rows = [["scenariusz", req.name]] if req.name else []
    return rows + [[k, v] for k, v in full.items()]


def build_csv(meta: list[list], params: list[list], rows: list[list]) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf, lineterminator="\n")
    w.writerow([META_TITLE])
    for r in meta:
        w.writerow([safe(c) for c in r])
    w.writerow([])
    w.writerow(["# Parametry (mld zł dla wyników)"])
    w.writerow(PARAM_HEADER)
    for r in params:
        w.writerow([safe(c) for c in r])
    w.writerow([])
    w.writerow(["# Wyniki"])
    w.writerow(HEADER)
    for r in rows:
        w.writerow([safe(c) for c in r])
    return ("﻿" + buf.getvalue()).encode("utf-8")


def build_xlsx(meta: list[list], params: list[list], rows: list[list]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Wyniki"
    ws.append(HEADER)
    for r in rows:
        ws.append([safe(c) for c in r])
    wp = wb.create_sheet("Parametry")
    wp.append(PARAM_HEADER)
    for r in params:
        wp.append([safe(c) for c in r])
    wm = wb.create_sheet("Metadane")
    for r in meta:
        wm.append([safe(c) for c in r])
    out = io.BytesIO()
    wb.save(out)
    wb.close()
    return out.getvalue()


router = APIRouter(prefix="/export", tags=["export"])


@router.post("")
def post_export(request: Request, req: ExportRequest, format: Literal["csv", "xlsx"] = Query(...)) -> Response:
    declared = request.headers.get("content-length", "")
    if declared.isdigit() and int(declared) > MAX_BODY_BYTES:
        raise HTTPException(status_code=413, detail="Żądanie zbyt duże (maks. 1 MB).")
    rows = result_rows(req.result)
    params = param_rows(req)
    meta = meta_rows(req.result)
    stamp = datetime.now(UTC).strftime("%Y%m%d")
    if format == "csv":
        content, media = build_csv(meta, params, rows), "text/csv; charset=utf-8"
    else:
        content, media = build_xlsx(meta, params, rows), XLSX_TYPE
    return Response(content, media_type=media, headers={"Content-Disposition": f'attachment; filename="symulacja-{stamp}.{format}"'})
