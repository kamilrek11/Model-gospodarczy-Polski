"""I-3 + C-1: exports carry provenance (data date, model version, synthetic flag, placeholder assumptions)."""

import csv
import io

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.api.export import ASSUMPTION_SENTENCE, SYNTHETIC_NOTE
from app.api.refresh import get_store
from app.main import app
from app.model.params import defaults
from app.model.simulate import simulate
from app.store import Store
from tests.test_simulate import synthetic_history, synthetic_model


@pytest.fixture()
def c(tmp_path):
    app.dependency_overrides[get_store] = lambda: Store(str(tmp_path / "t.db"))
    yield TestClient(app)
    app.dependency_overrides.clear()


@pytest.fixture(scope="module")
def result():
    hist = synthetic_history()
    res = simulate(synthetic_model(hist), hist, defaults(), horizon=2, n_paths=100, seed=1).to_dict()
    return {**res, "data_updated_at": "2026-01-02T03:04:05+00:00", "model_version": "abcdef123456", "synthetic": True}


def csv_rows(r) -> list[list[str]]:
    assert r.status_code == 200, r.text
    text = r.content.decode("utf-8")
    assert text.startswith("﻿")
    return list(csv.reader(io.StringIO(text[1:])))


def meta_of(rows: list[list[str]]) -> dict[str, str]:
    start = rows.index(["# Metadane"])
    end = rows.index(["# Parametry (mld zł dla wyników)"])
    return {r[0]: r[1] for r in rows[start + 1 : end] if len(r) >= 2}


def test_csv_starts_with_metadata_block(c, result):
    rows = csv_rows(c.post("/export?format=csv", json={"params": {}, "result": result}))
    assert rows[0] == ["# Metadane"]
    meta = meta_of(rows)
    assert meta["dane_z"] == "2026-01-02T03:04:05+00:00"
    assert meta["wersja_modelu"] == "abcdef123456"
    assert meta["dane_syntetyczne"] == "TAK"
    assert meta["uwaga"] == SYNTHETIC_NOTE and "DANE SYNTETYCZNE" in SYNTHETIC_NOTE
    assert meta["wygenerowano"].endswith("+00:00")
    placeholders = [a["key"] for a in result["assumptions"] if a["status"] == "assumption"]
    assert placeholders and meta["zalozenia_niezweryfikowane"] == ", ".join(placeholders)
    assert meta["zastrzezenie"].startswith(ASSUMPTION_SENTENCE) and "niezweryfikowanych" in meta["zastrzezenie"]
    assert all(k in meta["zastrzezenie"] for k in placeholders)
    assert "PLACEHOLDER" not in meta["zastrzezenie"]
    # parameters and results blocks follow, unchanged
    assert ["parametr", "wartość"] in rows and ["year", "tax", "median", "p10", "p90", "p2_5", "p97_5", "baseline"] in rows


def test_csv_real_data_has_no_synthetic_note(c, result):
    rows = csv_rows(c.post("/export?format=csv", json={"params": {}, "result": {**result, "synthetic": False}}))
    meta = meta_of(rows)
    assert meta["dane_syntetyczne"] == "NIE" and "uwaga" not in meta


def test_csv_missing_provenance_fields_tolerated(c, result):
    bare = {k: v for k, v in result.items() if k not in {"data_updated_at", "model_version", "synthetic", "assumptions"}}
    meta = meta_of(csv_rows(c.post("/export?format=csv", json={"params": {}, "result": bare})))
    assert meta["dane_z"] == "brak" and meta["wersja_modelu"] == "brak"
    assert meta["dane_syntetyczne"] == "nieznane"
    assert meta["zalozenia_niezweryfikowane"] == "brak informacji"


def test_xlsx_has_metadata_sheet(c, result):
    r = c.post("/export?format=xlsx", json={"params": {}, "result": result})
    assert r.status_code == 200
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["Wyniki", "Parametry", "Metadane"]
    meta = {row[0]: row[1] for row in wb["Metadane"].iter_rows(values_only=True)}
    wb.close()
    assert meta["dane_z"] == "2026-01-02T03:04:05+00:00"
    assert meta["dane_syntetyczne"] == "TAK" and meta["uwaga"] == SYNTHETIC_NOTE
    assert meta["zastrzezenie"].startswith(ASSUMPTION_SENTENCE)
    assert "PIT_INCOME_DIST_DEFAULT" in meta["zalozenia_niezweryfikowane"]


def test_provenance_text_is_formula_safe(c, result):
    evil = {**result, "model_version": "=cmd()", "assumptions": [{"key": "@SUM(1)", "status": "assumption", "value": 1, "note": ""}]}
    text = c.post("/export?format=csv", json={"params": {}, "result": evil}).content.decode("utf-8-sig")
    assert "'=cmd()" in text and "'@SUM(1)" in text


@pytest.mark.parametrize(
    "bad",
    [
        {"synthetic": "yes"},
        {"data_updated_at": 5},
        {"model_version": "x" * 200},
        {"assumptions": "nope"},
        {"assumptions": [{"key": "A"}] * 500},
        {"assumptions": [{"status": "assumption"}]},
        {"assumptions": [{"key": "A", "status": "sourced", "source": "x" * 3000}]},
        {"assumptions": [{"key": "A", "status": "sourced", "year": "abc"}]},
    ],
)
def test_provenance_fields_validated(c, result, bad):
    r = c.post("/export?format=csv", json={"params": {}, "result": {**result, **bad}})
    assert r.status_code == 422, r.text


def _with(result, rows):
    return {**result, "assumptions": rows}


def _row(key, status, **kw):
    return {"key": key, "value": 1, "status": status, "note": "n", "source": "", "year": None, "url": "", **kw}


def test_disclaimer_lists_only_assumption_constants(c, result):
    rows = [_row("A_SRC", "sourced", source="MF", year=2021), _row("B_DATA", "from_data", year=2024), _row("C_ASM", "assumption"), _row("D_STRUCT", "model-structure")]
    meta = meta_of(csv_rows(c.post("/export?format=csv", json={"params": {}, "result": _with(result, rows)})))
    assert meta["zalozenia_niezweryfikowane"] == "C_ASM"
    assert "C_ASM" in meta["zastrzezenie"]
    for k in ("A_SRC", "B_DATA", "D_STRUCT"):
        assert k not in meta["zastrzezenie"]


def test_no_disclaimer_when_no_assumption_constants(c, result):
    rows = [_row("A_SRC", "sourced", source="MF", year=2021), _row("B_DATA", "from_data", year=2024)]
    meta = meta_of(csv_rows(c.post("/export?format=csv", json={"params": {}, "result": _with(result, rows)})))
    assert meta["zalozenia_niezweryfikowane"] == "brak"
    assert "zastrzezenie" not in meta
    assert "niezweryfikowan" not in str(meta).replace("zalozenia_niezweryfikowane", "")


def test_metadata_carries_status_source_year_per_constant(c, result):
    rows = [_row("A_SRC", "sourced", source="MF CIT-8", year=2021, url="https://x.test/a"), _row("B_DATA", "from_data", year=2024), _row("D_STRUCT", "model-structure")]
    meta = meta_of(csv_rows(c.post("/export?format=csv", json={"params": {}, "result": _with(result, rows)})))
    assert meta["stala_A_SRC"] == "ze źródła; MF CIT-8; 2021; https://x.test/a"
    assert meta["stala_B_DATA"].startswith("z danych") and "2024" in meta["stala_B_DATA"]
    assert "stala_D_STRUCT" not in meta


def test_xlsx_metadata_has_constant_rows(c, result):
    r = c.post("/export?format=xlsx", json={"params": {}, "result": result})
    wb = load_workbook(io.BytesIO(r.content))
    meta = {row[0]: row[1] for row in wb["Metadane"].iter_rows(values_only=True)}
    wb.close()
    assert meta["stala_VAT_GAP"].startswith("ze źródła")


def test_legacy_placeholder_status_counts_as_assumption(c, result):
    rows = [_row("OLD", "PLACEHOLDER-UNVERIFIED"), _row("NEW", "assumption"), _row("SRC", "sourced", source="MF", year=2021)]
    meta = meta_of(csv_rows(c.post("/export?format=csv", json={"params": {}, "result": {**result, "assumptions": rows}})))
    assert meta["zalozenia_niezweryfikowane"] == "OLD, NEW"
    assert meta["zastrzezenie"].startswith(ASSUMPTION_SENTENCE) and "OLD" in meta["zastrzezenie"]
    assert meta["stala_OLD"].startswith("założenie")
    assert "PLACEHOLDER" not in meta["zastrzezenie"]


def test_missing_assumptions_field_still_has_disclaimer(c, result):
    bare = {k: v for k, v in result.items() if k != "assumptions"}
    meta = meta_of(csv_rows(c.post("/export?format=csv", json={"params": {}, "result": bare})))
    assert meta["zastrzezenie"]


def test_constant_rows_carry_range_and_no_empty_fields(c, result):
    rows = [
        _row("NUM", "sourced", year=2021, low=0.05, high=0.12),
        _row("TXT", "assumption", range_text="standard 0,45–0,65"),
        _row("DATA", "from_data"),
    ]
    meta = meta_of(csv_rows(c.post("/export?format=csv", json={"params": {}, "result": _with(result, rows)})))
    assert meta["stala_NUM"] == "ze źródła; 2021; zakres 0,05–0,12"
    assert meta["stala_TXT"] == "założenie; standard 0,45–0,65"
    assert meta["stala_DATA"] == "z danych"


@pytest.mark.parametrize("bad", [{"low": "x"}, {"high": True}, {"range_text": 5}, {"range_text": "x" * 3000}])
def test_range_fields_validated(c, result, bad):
    r = c.post("/export?format=csv", json={"params": {}, "result": _with(result, [_row("A", "sourced", **bad)])})
    assert r.status_code == 422, r.text
