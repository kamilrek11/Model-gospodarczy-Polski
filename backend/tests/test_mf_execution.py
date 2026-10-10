import io
import json
import struct
import zipfile

import httpx
import pandas as pd
import pytest

from app.data.base import SourceError
from app.data.mf_execution import (
    RESOURCES_URL,
    MfBudgetExecutionClient,
    Resource,
    extract_revenue,
    find_member,
    list_resources,
)
from tests.mf_fixtures import LABEL_PIT, build_workbook, fixture_path

MEMBER_007 = "007- DOCHODY BP ZESTAWIENIE OG#U00d3#U0141EM.xlsx"
PUBLISHED = {
    2023: {"rev_vat": 244267385.71, "rev_cit": 67883082.34, "rev_pit": 91665807.67},
    2024: {"rev_vat": 287683368.91, "rev_cit": 60239719.55, "rev_pit": 97599844.06},
    2025: {"rev_vat": 321646381.78, "rev_cit": 63686737.92, "rev_pit": 28752115.01},
}


def fixture(year: int) -> bytes:
    return fixture_path(year).read_bytes()


def make_zip(members: dict[str, bytes], compression=zipfile.ZIP_DEFLATED) -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", compression) as z:
        for name, data in members.items():
            z.writestr(name, data)
    return buf.getvalue()


def report_zip(year: int, workbook: bytes | None = None) -> bytes:
    """Outer ZIP -> 'TOM I' ZIP (+ 'TOM II' decoy) -> 007 workbook, like the MF archives."""
    tom1 = make_zip({
        "XLS/006-TYTU#U0141-DOCHODY BUD#U017bETU PA#U0143STWA.XLS": b"title page",
        f"XLS/{MEMBER_007}": workbook if workbook is not None else fixture(year),
    })
    tom2 = make_zip({"07 - BUD#U017bETY WOJEWOD#U00d3W - WYDATKI.xlsx": b"not this one"})
    return make_zip({
        f"Sprawozdanie za {year} rok TOM II XLS.zip": tom2,
        f"Sprawozdanie za {year} rok TOM I XLS.zip": tom1,
    })


def resource_item(rid: str, year: int, link: str, size: int) -> dict:
    return {
        "id": rid,
        "type": "resource",
        "attributes": {
            "title": f"Część tabelaryczna sprawozdania z wykonania budżetu państwa w {year} r.",
            "link": link,
            "file_size": size,
            "data_date": f"{year + 1}-06-01",
            "format": "zip",
        },
    }


class FakeDaneGov:
    """MockTransport handler serving the resource list and per-year ZIPs; records every request."""

    def __init__(self, zips: dict[int, bytes], *, per_page_limit: int | None = None):
        self.zips = dict(zips)
        self.requests: list[httpx.URL] = []
        self.overrides: dict[str, httpx.Response] = {}
        self.items = [
            resource_item(str(1000 + y), y, f"https://api.dane.gov.pl/media/resources/{y}/report_{y}.zip", len(z))
            for y, z in sorted(self.zips.items(), reverse=True)
        ]
        self.per_page_limit = per_page_limit

    def __call__(self, request: httpx.Request) -> httpx.Response:
        self.requests.append(request.url)
        url = str(request.url)
        if url in self.overrides:
            return self.overrides[url]
        if request.url.path == "/1.4/datasets/163/resources":
            page = int(request.url.params.get("page", "1"))
            per_page = self.per_page_limit or int(request.url.params.get("per_page", "50"))
            chunk = self.items[(page - 1) * per_page: page * per_page]
            return httpx.Response(200, json={"data": chunk, "meta": {"count": len(self.items)}, "links": {}})
        for y, z in self.zips.items():
            if request.url.path == f"/media/resources/{y}/report_{y}.zip":
                return httpx.Response(200, content=z, headers={"Content-Type": "application/zip"})
        return httpx.Response(404)

    def downloads(self) -> list[str]:
        return [str(u) for u in self.requests if "/media/" in u.path]


def client_for(fake, tmp_path=None, **kw) -> MfBudgetExecutionClient:
    http = httpx.Client(transport=httpx.MockTransport(fake))
    cache = str(tmp_path / "cache") if tmp_path is not None else None
    return MfBudgetExecutionClient(http, cache, sleep=lambda s: None, **kw)


# --- extract_revenue -------------------------------------------------------------------------


def test_extract_revenue_2024_layout():
    assert extract_revenue(fixture(2024)) == PUBLISHED[2024]


def test_extract_revenue_2023_and_2025_layout():
    assert extract_revenue(fixture(2023)) == PUBLISHED[2023]
    assert extract_revenue(fixture(2025)) == PUBLISHED[2025]


def test_extract_revenue_2016_and_2018_layouts():
    # 2016 also has a second row numbered "1.4." (Podatek tonażowy): labels match in full
    assert extract_revenue(fixture(2016)) == {"rev_vat": 126584120.0, "rev_cit": 26381397.0, "rev_pit": 48232395.0}
    assert extract_revenue(fixture(2018)) == {"rev_vat": 174947070.50, "rev_cit": 34640852.72, "rev_pit": 59558738.49}


def test_extract_revenue_rounds_to_published_precision():
    wb = build_workbook(2024)
    # the real cells hold float noise from formulas, e.g. 287683368.91015995
    import openpyxl

    book = openpyxl.load_workbook(io.BytesIO(wb))
    ws = book.active
    for row in ws.iter_rows():
        if row[1].value == "1.1. Podatek od towarów i usług":
            row[3].value = 287683368.91015995
    buf = io.BytesIO()
    book.save(buf)
    assert extract_revenue(buf.getvalue())["rev_vat"] == 287683368.91


def test_extract_revenue_checks_year_in_title():
    assert extract_revenue(fixture(2024), year=2024) == PUBLISHED[2024]
    with pytest.raises(SourceError, match="2023"):
        extract_revenue(fixture(2024), year=2023)


def test_missing_row_label_raises():
    with pytest.raises(SourceError) as e:
        extract_revenue(build_workbook(2024, drop_labels=(LABEL_PIT,)))
    assert e.value.source == "mf_execution"
    assert "1.5. Podatek dochodowy od osób fizycznych" in e.value.message


def test_missing_unit_line_raises():
    with pytest.raises(SourceError, match="w tysiącach zł"):
        extract_revenue(build_workbook(2023, unit=None))


def test_execution_header_missing_falls_back_to_last_numeric_cell_with_note():
    notes: list[str] = []
    assert extract_revenue(build_workbook(2023, execution_header=None), notes=notes) == PUBLISHED[2023]
    assert notes and "last numeric" in notes[0]


def test_legacy_xls_rejected():
    ole2 = bytes.fromhex("d0cf11e0a1b11ae1") + b"\0" * 512
    with pytest.raises(SourceError, match="legacy .xls"):
        extract_revenue(ole2)


def test_garbage_workbook_raises_source_error():
    with pytest.raises(SourceError) as e:
        extract_revenue(b"PK\x03\x04 definitely not a workbook")
    assert e.value.source == "mf_execution"


# --- find_member ---------------------------------------------------------------------------


def test_find_member_handles_nested_zip_and_escaped_names():
    assert find_member(report_zip(2024)) == fixture(2024)
    # directly in the outer archive, UTF-8 name, other case/spacing
    direct = make_zip({"007-dochody bp zestawienie ogółem (str.1.1).XLSX": b"wb"})
    assert find_member(direct) == b"wb"


def test_find_member_depth_limit():
    deep = make_zip({"a TOM I.zip": make_zip({"b TOM I.zip": make_zip({"c TOM I.zip": make_zip({MEMBER_007: b"x"})})})})
    with pytest.raises(SourceError, match="not found"):
        find_member(deep)
    two_levels = make_zip({"a TOM I.zip": make_zip({"b TOM I.zip": make_zip({MEMBER_007: b"x"})})})
    assert find_member(two_levels) == b"x"


def test_find_member_ambiguous_raises():
    z = make_zip({MEMBER_007: b"a", "XLS/007-DOCHODY BP ZESTAWIENIE OGOLEM.xlsx": b"b"})
    with pytest.raises(SourceError, match="ambiguous"):
        find_member(z)


def _patch_sizes(data: bytes, name: str, size: int) -> bytes:
    """Rewrite the uncompressed size of `name` in the central directory (and local header)."""
    buf = bytearray(data)
    enc = name.encode()
    pos = 0
    while True:
        pos = buf.find(b"PK\x01\x02", pos)
        if pos < 0:
            break
        n = struct.unpack_from("<H", buf, pos + 28)[0]
        if bytes(buf[pos + 46: pos + 46 + n]) == enc:
            struct.pack_into("<I", buf, pos + 24, size)
            local = struct.unpack_from("<I", buf, pos + 42)[0]
            struct.pack_into("<I", buf, local + 22, size)
        pos += 4
    return bytes(buf)


def test_oversized_member_rejected(monkeypatch):
    z = _patch_sizes(make_zip({MEMBER_007: b"small"}), MEMBER_007, 500_000_000)
    assert zipfile.ZipFile(io.BytesIO(z)).getinfo(MEMBER_007).file_size == 500_000_000

    def no_read(*a, **k):
        raise AssertionError("an oversized member must be rejected before reading")

    monkeypatch.setattr(zipfile.ZipFile, "open", no_read)
    with pytest.raises(SourceError, match="500000000"):
        find_member(z, max_member_bytes=30_000_000)


def test_lying_member_header_enforced_while_reading():
    big = b"0" * 50_000
    z = _patch_sizes(make_zip({MEMBER_007: big}), MEMBER_007, 100)
    with pytest.raises(SourceError) as e:
        find_member(z, max_member_bytes=1_000)
    assert e.value.source == "mf_execution"


def test_too_many_members_rejected():
    z = make_zip({f"f{i}.txt": b"" for i in range(30)} | {MEMBER_007: b"x"})
    with pytest.raises(SourceError, match="members"):
        find_member(z, max_members=20)


def test_zip_slip_names_never_touch_disk(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    evil = make_zip({"../../x": b"evil", f"../../{MEMBER_007}": b"evil 007", f"/abs/{MEMBER_007}": b"abs"})
    with pytest.raises(SourceError, match="not found"):
        find_member(evil)
    mixed = make_zip({"../../x": b"evil", f"../../{MEMBER_007}": b"evil 007", f"XLS/{MEMBER_007}": b"good"})
    assert find_member(mixed) == b"good"
    assert list(tmp_path.iterdir()) == []
    assert not (tmp_path.parent.parent / "x").exists()


def test_not_a_zip_raises():
    with pytest.raises(SourceError) as e:
        find_member(b"<html>maintenance</html>")
    assert e.value.source == "mf_execution"


# --- list_resources ------------------------------------------------------------------------


def test_list_resources_paginates_and_parses_year_from_title():
    fake = FakeDaneGov({2023: b"a", 2024: b"b", 2025: b"c"}, per_page_limit=2)
    http = httpx.Client(transport=httpx.MockTransport(fake))
    res = list_resources(http)
    assert sorted(r.year for r in res) == [2023, 2024, 2025]
    r24 = next(r for r in res if r.year == 2024)
    assert r24 == Resource(year=2024, url="https://api.dane.gov.pl/media/resources/2024/report_2024.zip", id="3024", size=1)
    assert all(str(u).startswith(RESOURCES_URL) or "/media/" in u.path for u in fake.requests)
    assert len([u for u in fake.requests if u.path.endswith("/resources")]) == 2


def test_resource_urls_must_be_https_and_allowlisted():
    fake = FakeDaneGov({2024: b"b"})
    fake.items += [
        resource_item("1", 2019, "http://api.dane.gov.pl/media/resources/x.zip", 10),
        resource_item("2", 2020, "https://evil.example/media/resources/x.zip", 10),
        resource_item("3", 2021, "https://api.dane.gov.pl.evil.example/x.zip", 10),
        resource_item("4", 2022, "ftp://dane.gov.pl/x.zip", 10),
        resource_item("5", 2018, "https://www.dane.gov.pl/media/resources/2018.zip", 10),
    ]
    warnings: list[str] = []
    http = httpx.Client(transport=httpx.MockTransport(fake))
    res = list_resources(http, warnings=warnings)
    assert sorted(r.year for r in res) == [2018, 2024]
    assert len(warnings) == 4 and all("rejected" in w for w in warnings)
    assert any("2020" in w for w in warnings)


def test_list_resources_title_without_year_is_skipped_with_warning():
    fake = FakeDaneGov({2024: b"b"})
    item = resource_item("9", 2000, "https://api.dane.gov.pl/media/resources/a.zip", 1)
    item["attributes"]["title"] = "Część tabelaryczna sprawozdania – aktualizacja"
    fake.items.append(item)
    warnings: list[str] = []
    res = list_resources(httpx.Client(transport=httpx.MockTransport(fake)), warnings=warnings)
    assert [r.year for r in res] == [2024]
    assert "year" in warnings[0]


def test_list_resources_errors_are_source_errors():
    for resp in (httpx.Response(404), httpx.Response(200, text="not json"), httpx.Response(200, json={"x": 1})):
        http = httpx.Client(transport=httpx.MockTransport(lambda r, resp=resp: resp))
        with pytest.raises(SourceError) as e:
            list_resources(http, sleep=lambda s: None)
        assert e.value.source == "mf_execution"


# --- client --------------------------------------------------------------------------------


def test_fetch_returns_billion_zl_frames():
    fake = FakeDaneGov({y: report_zip(y) for y in (2023, 2024, 2025)})
    out = client_for(fake).fetch()
    assert set(out) == {"rev_vat", "rev_cit", "rev_pit"}
    for s, df in out.items():
        assert isinstance(df, pd.DataFrame)
        assert df.index.name == "year" and df.index.dtype == "int64"
        assert df.index.tolist() == [2023, 2024, 2025]
        assert df["value"].dtype == "float64"
    assert out["rev_pit"].loc[2025, "value"] == pytest.approx(28.75211501, rel=1e-12)


def test_values_converted_thousand_zl_to_billion():
    fake = FakeDaneGov({2024: report_zip(2024)})
    out = client_for(fake).fetch()
    assert out["rev_vat"].loc[2024, "value"] == pytest.approx(287.68336891, rel=1e-12)
    assert out["rev_cit"].loc[2024, "value"] == pytest.approx(60.23971955, rel=1e-12)


def test_years_filter_limits_downloads():
    fake = FakeDaneGov({y: report_zip(y) for y in (2023, 2024, 2025)})
    out = client_for(fake, years=range(2024, 2025)).fetch()
    assert out["rev_vat"].index.tolist() == [2024]
    assert len(fake.downloads()) == 1


def test_cache_hit_skips_download(tmp_path):
    fake = FakeDaneGov({y: report_zip(y) for y in (2023, 2024)})
    first = client_for(fake, tmp_path).fetch()
    assert len(fake.downloads()) == 2
    files = sorted(p.name for p in (tmp_path / "cache").iterdir())
    assert all(f.endswith(".json") for f in files) and len(files) == 2
    fake.requests.clear()
    second = client_for(fake, tmp_path).fetch()
    assert fake.downloads() == []
    for s in first:
        pd.testing.assert_frame_equal(first[s], second[s])
    cached = json.loads((tmp_path / "cache" / files[0]).read_text(encoding="utf-8"))
    assert set(cached) >= {"resource_id", "size", "values"}


def test_cache_invalidated_when_resource_changes(tmp_path):
    fake = FakeDaneGov({2024: report_zip(2024)})
    client_for(fake, tmp_path).fetch()
    fake.items[0]["attributes"]["file_size"] += 1  # re-published resource
    fake.requests.clear()
    client_for(fake, tmp_path).fetch()
    assert len(fake.downloads()) == 1


def test_corrupted_cache_ignored_and_refetched(tmp_path):
    fake = FakeDaneGov({2024: report_zip(2024)})
    client_for(fake, tmp_path).fetch()
    (path,) = list((tmp_path / "cache").iterdir())
    path.write_text("{not json", encoding="utf-8")
    fake.requests.clear()
    out = client_for(fake, tmp_path).fetch()
    assert len(fake.downloads()) == 1
    assert out["rev_vat"].loc[2024, "value"] == pytest.approx(287.68336891)
    assert json.loads(path.read_text(encoding="utf-8"))["values"]["rev_vat"] == 287683368.91


def test_year_failure_reports_year_in_source_error_and_other_years_still_returned_policy():
    fake = FakeDaneGov({2023: report_zip(2023), 2024: report_zip(2024, build_workbook(2024, drop_labels=(LABEL_PIT,)))})
    client = client_for(fake)
    out = client.fetch()
    assert out["rev_vat"].index.tolist() == [2023]
    assert len(client.warnings) == 1 and "2024" in client.warnings[0] and LABEL_PIT in client.warnings[0]

    bad = FakeDaneGov({2024: report_zip(2024, build_workbook(2024, drop_labels=(LABEL_PIT,))), 2025: b"not a zip"})
    client = client_for(bad)
    with pytest.raises(SourceError) as e:
        client.fetch()
    assert e.value.source == "mf_execution"
    assert "2024" in e.value.message and "2025" in e.value.message


def test_unsupported_year_is_cached_as_unsupported(tmp_path):
    ole2 = bytes.fromhex("d0cf11e0a1b11ae1") + b"\0" * 512
    tom1 = make_zip({"007- DOCHODY BP ZESTAWIENIE OGOLEM (STR. 1.1).XLS": ole2})
    fake = FakeDaneGov({2015: make_zip({"x TOM I_XLS.zip": tom1}), 2024: report_zip(2024)})
    client = client_for(fake, tmp_path)
    client.fetch()
    assert any("2015" in w and "legacy .xls" in w for w in client.warnings)
    fake.requests.clear()
    again = client_for(fake, tmp_path)
    again.fetch()
    assert fake.downloads() == []
    assert any("2015" in w and "legacy .xls" in w for w in again.warnings)


def test_download_limit_enforced_while_streaming():
    pulled = []

    class Body(httpx.SyncByteStream):
        def __iter__(self):
            for _ in range(1000):
                pulled.append(1)
                yield b"0" * 10_000

    fake = FakeDaneGov({2023: report_zip(2023), 2024: b"x"})
    url = "https://api.dane.gov.pl/media/resources/2024/report_2024.zip"
    fake.overrides[url] = httpx.Response(200, stream=Body())
    fake.items[0]["attributes"]["file_size"] = 1  # the API under-reports the size
    client = client_for(fake, max_download_bytes=50_000)
    out = client.fetch()
    assert out["rev_vat"].index.tolist() == [2023]
    assert any("2024" in w and "50000" in w for w in client.warnings)
    assert len(pulled) < 10


def test_declared_size_over_limit_skips_download():
    fake = FakeDaneGov({2023: report_zip(2023), 2024: report_zip(2024)})
    fake.items[0]["attributes"]["file_size"] = 10**9
    client = client_for(fake)
    client.fetch()
    assert len(fake.downloads()) == 1
    assert any("2024" in w for w in client.warnings)


def test_content_length_over_limit_rejected_before_body():
    fake = FakeDaneGov({2023: report_zip(2023), 2024: b"x"})
    url = "https://api.dane.gov.pl/media/resources/2024/report_2024.zip"
    fake.overrides[url] = httpx.Response(200, content=b"0" * 60_000)
    client = client_for(fake, max_download_bytes=50_000)
    client.fetch()
    assert any("2024" in w and "50000" in w for w in client.warnings)


def test_redirect_within_allowlist_is_followed_and_to_other_host_rejected():
    fake = FakeDaneGov({2023: report_zip(2023), 2024: report_zip(2024)})
    u23 = "https://api.dane.gov.pl/media/resources/2023/report_2023.zip"
    u24 = "https://api.dane.gov.pl/media/resources/2024/report_2024.zip"
    fake.overrides[u23] = httpx.Response(301, headers={"Location": "https://dane.gov.pl/media/resources/2023/report_2023.zip"})
    fake.overrides[u24] = httpx.Response(302, headers={"Location": "https://evil.example/media/resources/2024/report_2024.zip"})
    client = client_for(fake)
    out = client.fetch()
    assert out["rev_vat"].index.tolist() == [2023]
    assert any("2024" in w and "evil.example" in w for w in client.warnings)
    assert all(u.host != "evil.example" for u in fake.requests)


def test_retries_on_5xx_and_429_only():
    fake = FakeDaneGov({2024: report_zip(2024)})
    calls = {"n": 0}
    inner = fake.__call__

    def flaky(request):
        if "/media/" in request.url.path and calls["n"] < 2:
            calls["n"] += 1
            return httpx.Response(503 if calls["n"] == 1 else 429, headers={"Retry-After": "1"})
        return inner(request)

    sleeps = []
    http = httpx.Client(transport=httpx.MockTransport(flaky))
    out = MfBudgetExecutionClient(http, None, sleep=sleeps.append).fetch()
    assert out["rev_vat"].index.tolist() == [2024]
    assert sleeps == [1.0, 1.0]

    fake404 = FakeDaneGov({2024: report_zip(2024)})
    fake404.overrides["https://api.dane.gov.pl/media/resources/2024/report_2024.zip"] = httpx.Response(404)
    client = client_for(fake404)
    with pytest.raises(SourceError, match="404"):
        client.fetch()
    assert len(fake404.downloads()) == 1


def test_network_error_is_source_error():
    def boom(request):
        raise httpx.ConnectError("offline", request=request)

    http = httpx.Client(transport=httpx.MockTransport(boom))
    with pytest.raises(SourceError) as e:
        MfBudgetExecutionClient(http, None, sleep=lambda s: None).fetch()
    assert e.value.source == "mf_execution"


def test_client_name_and_constructor_rejects_bad_limits():
    http = httpx.Client(transport=httpx.MockTransport(lambda r: httpx.Response(404)))
    assert MfBudgetExecutionClient(http, None).name == "mf_execution"
    with pytest.raises(ValueError):
        MfBudgetExecutionClient(http, None, max_download_bytes=0)


def test_fixtures_are_small_reconstructions():
    for y in (2016, 2018, 2023, 2024, 2025):
        p = fixture_path(y)
        assert p.stat().st_size < 20_000
        assert zipfile.is_zipfile(p)


def test_fixtures_replicate_column_layouts():
    import openpyxl

    def exec_col(year):
        wb = openpyxl.load_workbook(io.BytesIO(fixture(year)), read_only=True)
        try:
            for row in wb.active.iter_rows():
                for c in row:
                    if c.value == "Wykonanie":
                        return c.column
        finally:
            wb.close()

    assert {y: exec_col(y) for y in (2016, 2018, 2023, 2024, 2025)} == {2016: 5, 2018: 3, 2023: 5, 2024: 4, 2025: 5}


def test_oversized_nested_archive_skipped_and_reported():
    stored = zipfile.ZIP_STORED
    inner = make_zip({"x TOM I_XLS.zip": make_zip({MEMBER_007: b"x", "pad.bin": b"0" * 5000}, stored)}, stored)
    big_pdfless = make_zip({"wrapper Tom I i II.zip": inner}, stored)
    with pytest.raises(SourceError, match="oversized") as e:
        find_member(big_pdfless, max_member_bytes=2000)
    assert "wrapper Tom I i II.zip" in e.value.message
    ok = make_zip({"a big TOM II.zip": b"0" * 5000, "b TOM I.zip": make_zip({MEMBER_007: b"x"})}, compression=zipfile.ZIP_STORED)
    assert find_member(ok, max_member_bytes=2000) == b"x"


def test_content_failure_cached_as_unsupported(tmp_path):
    fake = FakeDaneGov({2023: report_zip(2023), 2024: make_zip({"nothing.txt": b""})})
    client = client_for(fake, tmp_path)
    client.fetch()
    assert any("2024" in w and "not found" in w for w in client.warnings)
    fake.requests.clear()
    client_for(fake, tmp_path).fetch()
    assert fake.downloads() == []


def test_known_unsupported_resource_not_downloaded():
    from app.data.mf_execution import KNOWN_UNSUPPORTED_RESOURCES

    assert "19451" in KNOWN_UNSUPPORTED_RESOURCES  # the 2015 report (legacy .xls)
    fake = FakeDaneGov({2024: report_zip(2024)})
    fake.items.append(resource_item("19451", 2015, "https://api.dane.gov.pl/media/resources/2015/report_2015.zip", 38428339))
    client = client_for(fake)
    out = client.fetch()
    assert out["rev_vat"].index.tolist() == [2024]
    assert all("2015" not in d for d in fake.downloads())
    assert any(w.startswith("2015:") and "legacy .xls" in w for w in client.warnings)


# --- fix round 1: decompression budget, archive cap, cache policy -----------------------------


def _counting_open(monkeypatch) -> dict:
    """Count every decompressed byte returned by ZipFile.open(...).read()."""
    counter = {"bytes": 0}
    real_open = zipfile.ZipFile.open

    def open_(self, *a, **k):
        f = real_open(self, *a, **k)
        real_read = f.read

        def read(n=-1):
            data = real_read(n)
            counter["bytes"] += len(data)
            return data

        f.read = read
        return f

    monkeypatch.setattr(zipfile.ZipFile, "open", open_)
    return counter


def _amplifier(n_mid: int, k_leaf: int, leaf_bytes: int) -> bytes:
    """outer (deflated) -> n_mid deflated 'TOM I' zips -> k_leaf stored 'TOM I' zips of leaf_bytes zeros."""
    leaf = make_zip({"pad.bin": b"\0" * leaf_bytes}, zipfile.ZIP_STORED)
    mid = make_zip({f"leaf{j} TOM I XLS.zip": leaf for j in range(k_leaf)})
    return make_zip({f"mid{i} TOM I XLS.zip": mid for i in range(n_mid)})


def test_nested_amplification_stops_at_total_budget(monkeypatch):
    bomb = _amplifier(n_mid=3, k_leaf=3, leaf_bytes=90_000)  # each archive below the member limit
    assert len(bomb) < 20_000
    counter = _counting_open(monkeypatch)
    with pytest.raises(SourceError, match="decompression budget") as e:
        find_member(bomb, max_member_bytes=1_000_000, max_total_decompressed_bytes=400_000)
    assert e.value.source == "mf_execution"
    assert counter["bytes"] <= 400_000 + 1_000_000  # bounded: at most one chunk past the budget


def test_nested_amplification_stops_at_archive_cap(monkeypatch):
    bomb = _amplifier(n_mid=6, k_leaf=6, leaf_bytes=1_000)
    counter = _counting_open(monkeypatch)
    with pytest.raises(SourceError, match="nested archives") as e:
        find_member(bomb, max_archives=10)
    assert "10" in e.value.message
    assert counter["bytes"] < 200_000


def test_only_tom_i_archives_are_opened():
    z = make_zip({"Sprawozdanie TOM II XLS.zip": make_zip({MEMBER_007: b"wrong volume"}),
                  "inne.zip": make_zip({MEMBER_007: b"other"})})
    with pytest.raises(SourceError, match="not found"):
        find_member(z)


def test_html_page_of_wrong_length_not_cached_as_unsupported(tmp_path):
    good = report_zip(2024)
    fake = FakeDaneGov({2024: good})
    url = "https://api.dane.gov.pl/media/resources/2024/report_2024.zip"
    fake.overrides[url] = httpx.Response(200, html="<html>przerwa techniczna</html>")
    client = client_for(fake, tmp_path)
    with pytest.raises(SourceError, match="2024"):
        client.fetch()
    del fake.overrides[url]
    fake.requests.clear()
    out = client_for(fake, tmp_path).fetch()
    assert len(fake.downloads()) == 1
    assert out["rev_vat"].loc[2024, "value"] == pytest.approx(287.68336891)


def test_limit_dependent_failure_not_cached(tmp_path):
    fake = FakeDaneGov({2023: report_zip(2023), 2024: report_zip(2024)})
    client = client_for(fake, tmp_path, max_member_bytes=1_000)
    with pytest.raises(SourceError):
        client.fetch()
    fake.requests.clear()
    out = client_for(fake, tmp_path).fetch()
    assert len(fake.downloads()) == 2
    assert out["rev_vat"].index.tolist() == [2023, 2024]


def test_changing_a_limit_invalidates_cached_entry(tmp_path):
    fake = FakeDaneGov({2024: report_zip(2024)})
    client_for(fake, tmp_path).fetch()
    fake.requests.clear()
    client_for(fake, tmp_path).fetch()
    assert fake.downloads() == []
    client_for(fake, tmp_path, max_member_bytes=29_000_000).fetch()
    assert len(fake.downloads()) == 1
