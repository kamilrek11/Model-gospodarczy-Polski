"""Reconstructed MF '007 - DOCHODY BP ZESTAWIENIE OGÓŁEM' workbooks (tests only).

The layouts (sheet name, title/header/unit rows, label/plan/execution columns, the row order and
the extra column-A short labels of some years) replicate the real files of the MF report
"Część tabelaryczna sprawozdania z wykonania budżetu państwa" (dane.gov.pl dataset 163) as
inspected on 2026-10-10; the numbers are the published report values (tys. zł, 2 decimals).
They are not copies of the MF files. `python -m tests.mf_fixtures` (from backend/) rewrites
tests/fixtures/real/mf_007_*.xlsx.
"""

import io
from pathlib import Path

import openpyxl

UNIT = " w tysiącach zł"
PLAN_HEADER = "Budżet \nwedług ustawy budżetowej"
LABEL_VAT = "1.1. Podatek od towarów i usług"
LABEL_CIT = "1.4. Podatek dochodowy od osób prawnych"
LABEL_PIT = "1.5. Podatek dochodowy od osób fizycznych"

# layout: sheet, title cell, header row, columns (short label or None, numbered label, plan, execution)
LAYOUTS = {
    # 2016, 2019-2023, 2025: short label in A, numbered label in C, plan D, execution E
    "AC_D_E": {"short": 1, "label": 3, "plan": 4, "exec": 5, "title_col": 2, "header_label_col": 2},
    # 2017, 2018: label in A, plan B, execution C
    "A_B_C": {"short": None, "label": 1, "plan": 2, "exec": 3, "title_col": 1, "header_label_col": 1},
    # 2024: label in B, plan C, execution D
    "B_C_D": {"short": None, "label": 2, "plan": 3, "exec": 4, "title_col": 1, "header_label_col": 1},
}

# year -> (layout, sheet name, title row, header row, rows); rows = (short, numbered, plan, execution)
YEARS = {
    2016: ("AC_D_E", "Zał 1", 2, 6, [
        ("DOCHODY OGÓŁEM", "DOCHODY OGÓŁEM", 313808526, 314683570),
        None,
        ("Dochody podatkowe", "1. Dochody podatkowe", 276140000, 273138413),
        None,
        ("Podatek od towarów i usług", LABEL_VAT, 128683000, 126584120),
        ("Podatek akcyzowy", "1.2. Podatek akcyzowy", 64083000, 65749274),
        ("Podatek od gier", "1.3. Podatek od gier", 1383000, 1406925),
        ("Podatek dochodowy od osób prawnych", LABEL_CIT, 26067000, 26381397),
        ("Podatek dochodowy od osób fizycznych", LABEL_PIT, 46894000, 48232395),
        # the real 2016 file numbers this row "1.4." too: labels must match in full
        ("Podatek tonażowy", "1.4. Podatek tonażowy", 0, None),
        ("Podatek od wydobycia niektórych kopalin", "1.6. Podatek od wydobycia niektórych kopalin", 1530000, 1277488),
    ]),
    2018: ("A_B_C", "ZAŁ 1", 1, 5, [
        (None, "DOCHODY OGÓŁEM", 355705405, 380048139.61),
        None,
        (None, "1. Dochody podatkowe", 331672637, 349353842.85),
        None,
        (None, LABEL_VAT, 166000000, 174947070.50),
        (None, "1.2. Podatek akcyzowy", 70000000, 72108486.17),
        (None, "1.3. Podatek od gier", 1913982, 1901914.94),
        (None, LABEL_CIT, 32400000, 34640852.72),
        (None, LABEL_PIT, 55500000, 59558738.49),
        (None, "1.6. Podatek od wydobycia niektórych kopalin", 1290000, 1689121.09),
    ]),
    2023: ("AC_D_E", "2023 Załacznik nr 1", 3, 5, [
        ("DOCHODY OGÓŁEM", "DOCHODY OGÓŁEM", 601378366, 573957931.78),
        None,
        ("Dochody podatkowe", "1. Dochody podatkowe", 536833242, 506866722.69),
        None,
        ("Podatek od towarów i usług", LABEL_VAT, 272900000, 244267385.71),
        ("Podatek akcyzowy", "1.2. Podatek akcyzowy", 83800000, 84798843.54),
        ("Podatek od gier", "1.3. Podatek od gier", 4500000, 4475681.87),
        ("Podatek dochodowy od osób prawnych", LABEL_CIT, 77606988, 67883082.34),
        ("Podatek dochodowy od osób fizycznych", LABEL_PIT, 83570254, 91665807.67),
        ("Podatek od wydobycia niektórych kopalin", "1.6. Podatek od wydobycia niektórych kopalin", 4600000, 3912329.40),
    ]),
    2024: ("B_C_D", "2024 Załacznik nr 1", 3, 5, [
        (None, "DOCHODY OGÓŁEM", 626075747, 623239830.63),
        None,
        (None, "1. Dochody podatkowe", 559605826, 555936540.32),
        None,
        (None, LABEL_VAT, 293500000, 287683368.91),
        (None, "1.2. Podatek akcyzowy", 90311000, 90315131.19),
        (None, "1.3. Podatek od gier", 5000000, 5266225.62),
        (None, LABEL_CIT, 58556935, 60239719.55),
        (None, LABEL_PIT, 97745431, 97599844.06),
        (None, "1.6. Podatek od wydobycia niektórych kopalin", 3900000, 4034085.79),
    ]),
    2025: ("AC_D_E", "2025 Załacznik nr 1", 3, 5, [
        ("DOCHODY OGÓŁEM", "DOCHODY OGÓŁEM", 632848215, 594524379.52),
        None,
        ("Dochody podatkowe", "1. Dochody podatkowe", 570772190, 529224947.93),
        None,
        ("Podatek od towarów i usług", LABEL_VAT, 349526000, 321646381.78),
        ("Podatek akcyzowy", "1.2. Podatek akcyzowy", 98136000, 92519250.19),
        ("Podatek od gier", "1.3. Podatek od gier", 5300000, 6193261.34),
        ("Podatek dochodowy od osób prawnych", LABEL_CIT, 70718957, 63686737.92),
        ("Podatek dochodowy od osób fizycznych", LABEL_PIT, 31735610, 28752115.01),
        ("Podatek od wydobycia niektórych kopalin", "1.6. Podatek od wydobycia niektórych kopalin", 3500000, 4750531.93),
    ]),
}

FIXTURE_DIR = Path(__file__).parent / "fixtures" / "real"


def build_workbook(
    year: int,
    *,
    drop_labels: tuple[str, ...] = (),
    execution_header: str | None = "Wykonanie",
    unit: str | None = UNIT,
    title_year: int | None = None,
) -> bytes:
    """Workbook bytes in the layout of `year` (variants for error tests via the keyword args)."""
    layout_name, sheet, title_row, header_row, rows = YEARS[year]
    lay = LAYOUTS[layout_name]
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet
    ws.cell(title_row, lay["title_col"], f"DOCHODY BUDŻETU PAŃSTWA W {title_year or year} R.")
    ws.cell(header_row, lay["header_label_col"], "Wyszczególnienie")
    ws.cell(header_row, lay["plan"], PLAN_HEADER)
    if execution_header is not None:
        ws.cell(header_row, lay["exec"], execution_header)
    if unit is not None:
        ws.cell(header_row + 1, lay["plan"], unit)
    for n, col in enumerate((lay["label"], lay["plan"], lay["exec"]), 1):
        ws.cell(header_row + 2, col, n)
    r = header_row + 4
    for row in rows:
        if row is not None:
            short, label, plan, execution = row
            if label not in drop_labels:
                if short is not None and lay["short"] is not None:
                    ws.cell(r, lay["short"], short)
                ws.cell(r, lay["label"], label)
                ws.cell(r, lay["plan"], plan)
                if execution is not None:
                    ws.cell(r, lay["exec"], execution)
        r += 1
    buf = io.BytesIO()
    wb.save(buf)
    wb.close()
    return buf.getvalue()


def fixture_path(year: int) -> Path:
    return FIXTURE_DIR / f"mf_007_{year}.xlsx"


def write_fixtures() -> None:
    for year in YEARS:
        fixture_path(year).write_bytes(build_workbook(year))


if __name__ == "__main__":
    write_fixtures()
