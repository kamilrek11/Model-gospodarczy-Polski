import io

import numpy as np
import pandas as pd
import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.api.refresh import get_store
from app.main import app
from app.model.backtest import backtest
from app.model.macro import MACRO_VARS
from app.model.params import defaults
from app.model.simulate import simulate
from app.store import Store
from tests.synthetic import synthetic_levels
from tests.test_simulate import synthetic_history, synthetic_model

# Everything below is synthetic and lives only in pytest's tmp_path stores / memory.


def perfect_history(n_years: int = 24) -> dict[str, pd.DataFrame]:
    """Constant growth rates (zero VAR residuals); CIT and VAT revenue exactly proportional to their bases."""
    idx = pd.Index(range(2000, 2000 + n_years), dtype="int64", name="year")
    t = np.arange(n_years)
    growth = {"pkb_real": 0.02, "employment": 0.005, "wage_nominal": 0.04, "cpi": 0.02, "consumption_nominal": 0.05}
    hist = {v: pd.DataFrame({"value": 100.0 * np.exp(g * t)}, index=idx) for v, g in growth.items()}
    nominal_gdp = hist["pkb_real"]["value"] * hist["cpi"]["value"]
    gs = 0.4 * nominal_gdp
    hist["gross_surplus"] = pd.DataFrame({"value": gs}, index=idx)
    hist["rev_cit"] = pd.DataFrame({"value": 0.03 * gs}, index=idx)
    hist["rev_vat"] = pd.DataFrame({"value": 0.12 * hist["consumption_nominal"]["value"]}, index=idx)
    return hist


def test_backtest_perfect_series_zero_mape():
    res = backtest(perfect_history(), horizon=3, min_train=15)
    assert set(res.mape) == {"cit", "vat"}
    for tax, m in res.mape.items():
        assert m == pytest.approx(0.0, abs=1e-6), tax
    assert res.message is None


def test_backtest_coverage_between_0_and_1():
    res = backtest(synthetic_history(22), horizon=2, min_train=15)
    assert set(res.mape) == {"pit", "cit", "vat"}
    for d in (res.coverage80, res.coverage95):
        assert set(d) == set(res.mape)
        assert all(0.0 <= v <= 1.0 for v in d.values())
    assert all(v >= 0 for v in res.mape.values())
    assert any("obecnym prawie" in n for n in res.notes)
    assert backtest(synthetic_history(22), horizon=2, min_train=15) == res  # seeded


def test_backtest_not_enough_data_returns_message():
    res = backtest(synthetic_history(14), horizon=3)
    assert res.mape == {} and res.coverage80 == {} and res.coverage95 == {}
    assert res.message


def test_backtest_uses_realised_years_only():
    hist = synthetic_history(20)
    hist["rev_pit"] = hist["rev_pit"].iloc[:0]  # PIT has no data at all -> omitted, not NaN/0
    res = backtest(hist, horizon=2, min_train=15)
    assert "pit" not in res.mape and "pit" not in res.coverage80


@pytest.fixture()
def client(tmp_path):
    store = Store(str(tmp_path / "t.db"))
    app.dependency_overrides[get_store] = lambda: store
    yield TestClient(app), store
    app.dependency_overrides.clear()


def fill(store: Store, n_years: int = 20):
    hist = synthetic_history(n_years)
    model = synthetic_model(hist)
    store.save_snapshot({k: ("test", v) for k, v in hist.items()}, model.to_dict(), {"updated_at": "x"})
    return hist, model


def test_history_409_on_empty_store(client):
    c, _ = client
    r = c.get("/history")
    assert r.status_code == 409 and r.json()["detail"] == "brak danych – uruchom odświeżenie"
    assert c.get("/backtest").status_code == 409


def test_history_returns_revenues(client):
    c, store = client
    hist, _ = fill(store)
    body = c.get("/history").json()
    assert set(body) == {"pit", "cit", "vat", "synthetic"}
    assert body["synthetic"] is False
    assert len(body["pit"]) == 20
    assert body["vat"][0]["year"] == 2000
    assert body["vat"][0]["value"] == pytest.approx(hist["rev_vat"]["value"].iloc[0])


def test_backtest_endpoint(client):
    c, store = client
    fill(store, 22)
    r = c.get("/backtest")
    assert r.status_code == 200
    body = r.json()
    assert set(body) >= {"mape", "coverage80", "coverage95", "notes", "message"}


def test_backtest_endpoint_empty_message(client):
    c, store = client
    fill(store, 15)
    body = c.get("/backtest").json()
    assert body["mape"] == {} and body["message"]


def sim_result(hist, model, params=None):
    return simulate(model, hist, params or defaults(), horizon=3, n_paths=200, seed=1).to_dict()


def test_export_csv_contains_params_and_years(client):
    c, store = client
    hist, model = fill(store)
    result = sim_result(hist, model)
    r = c.post("/export?format=csv", json={"params": {"vat_rate_standard": 0.25}, "result": result})
    assert r.status_code == 200
    assert r.headers["content-type"].startswith("text/csv")
    assert "attachment" in r.headers["content-disposition"] and ".csv" in r.headers["content-disposition"]
    text = r.content.decode("utf-8-sig")
    assert "vat_rate_standard" in text and "0.25" in text
    assert "year,tax,median,p10,p90,p2_5,p97_5,baseline" in text
    for y in result["years"]:
        assert f"{y},vat," in text


def test_export_csv_escapes_formula_injection(client):
    c, store = client
    hist, model = fill(store)
    r = c.post("/export?format=csv", json={"params": {}, "result": sim_result(hist, model), "name": "=HYPERLINK(\"x\")"})
    assert r.status_code == 200
    text = r.content.decode("utf-8-sig")
    assert "'=HYPERLINK" in text and "\n=HYPERLINK" not in text and ",=HYPERLINK" not in text


def test_export_xlsx_has_data_sheets(client):
    c, store = client
    hist, model = fill(store)
    r = c.post("/export?format=xlsx", json={"params": {}, "result": sim_result(hist, model)})
    assert r.status_code == 200
    assert "spreadsheetml" in r.headers["content-type"]
    assert ".xlsx" in r.headers["content-disposition"]
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["Wyniki", "Parametry", "Metadane"]
    assert wb["Wyniki"].max_row > 3 and wb["Parametry"].max_row > 3
    wb.close()


def test_export_unknown_format_422(client):
    c, store = client
    hist, model = fill(store)
    r = c.post("/export?format=pdf", json={"params": {}, "result": sim_result(hist, model)})
    assert r.status_code == 422


@pytest.mark.parametrize("body", [{}, {"params": {}}, {"params": {}, "result": {"years": [2030]}}, {"params": {"bogus": 1}, "result": {"years": [], "taxes": {}, "baseline": {}}}, {"params": [], "result": {}}])
def test_export_malformed_body_422(client, body):
    c, _ = client
    assert c.post("/export?format=csv", json=body).status_code == 422


def test_backtest_min_train_never_below_min_obs_so_nothing_is_skipped_spuriously():
    from app.model.macro import MIN_OBS

    default = backtest(synthetic_history(22), horizon=2)
    assert default.n_skipped == 0 and default.n_origins > 0
    assert not any("Pominięto" in n for n in default.notes)
    low = backtest(synthetic_history(22), horizon=2, min_train=12)  # clamped to MIN_OBS
    assert low == default and MIN_OBS == 15


def test_backtest_reports_skipped_origins_on_real_fit_failures(monkeypatch):
    import app.model.backtest as bt_mod

    real_fit, calls = bt_mod.fit_macro, {"n": 0}

    def flaky(*a, **k):
        calls["n"] += 1
        if calls["n"] <= 2:
            raise np.linalg.LinAlgError("singular matrix")
        return real_fit(*a, **k)

    monkeypatch.setattr(bt_mod, "fit_macro", flaky)
    res = backtest(synthetic_history(22), horizon=2)
    assert res.n_skipped == 2 and res.n_origins > 0
    assert any("Pominięto 2" in n for n in res.notes)


@pytest.mark.parametrize("fmt", ["csv", "xlsx"])
@pytest.mark.parametrize("value", [[1, 2], {"a": 1}, "0.2", True, float("nan"), float("inf"), None])
def test_export_bad_param_values_422(client, fmt, value):
    c, store = client
    hist, model = fill(store)
    import json

    raw = json.dumps({"params": {"vat_rate_standard": value}, "result": sim_result(hist, model)})  # nan/inf -> NaN/Infinity tokens
    r = c.post(f"/export?format={fmt}", content=raw, headers={"content-type": "application/json"})
    assert r.status_code == 422, (value, r.status_code)


def test_export_oversized_payload_rejected(client):
    c, store = client
    hist, model = fill(store)
    result = sim_result(hist, model)
    big = {**result, "years": list(range(2000, 2200))}
    r = c.post("/export?format=csv", json={"params": {}, "result": big})
    assert r.status_code == 422 and "zbyt duży" in r.json()["detail"]
    r = c.post("/export?format=csv", json={"params": {}, "result": result, "name": "x" * 1_100_000})
    assert r.status_code == 413 and "zbyt duże" in r.json()["detail"]


@pytest.mark.parametrize("prefix", ["=", "+", "-", "@", "\t", "\r"])
def test_export_csv_guards_all_formula_prefixes(client, prefix):
    c, store = client
    hist, model = fill(store)
    r = c.post("/export?format=csv", json={"params": {}, "result": sim_result(hist, model), "name": prefix + "cmd"})
    assert ("'" + prefix + "cmd") in r.content.decode("utf-8-sig")


def test_backtest_skips_pit_targets_from_break_year_and_notes_it():
    from app.data.dataset import assemble
    from app.model.backtest import NOTE_PIT_BREAK
    from tests.real_dataset import real_results

    data = {k: df for k, (_, df) in assemble(real_results()).items()}
    assert data["rev_pit"].index.max() == 2025
    res = backtest(data, horizon=2)
    assert "pit" in res.mape and res.message is None
    assert NOTE_PIT_BREAK in res.notes
    assert any(n.startswith("W backteście PIT liczony jest według zasad sprzed 2025 r. (bez udziału budżetu państwa)") for n in res.notes)
    assert NOTE_PIT_BREAK.startswith("PIT w budżecie państwa: zmiana zasad podziału od 2025 r.")

    # PIT 2025 is never scored: any value there gives exactly the same PIT metrics as no value at all
    without = dict(data)
    without["rev_pit"] = data["rev_pit"].loc[:2024]
    ref = backtest(without, horizon=2)
    assert NOTE_PIT_BREAK not in ref.notes  # nothing skipped -> no note
    assert not any("zasad sprzed 2025" in n for n in ref.notes)
    for v in (1e-3, 1e4):
        odd = dict(data)
        pit = data["rev_pit"].copy()
        pit.loc[2025, "value"] = v
        odd["rev_pit"] = pit
        r = backtest(odd, horizon=2)
        assert (r.mape["pit"], r.coverage80["pit"], r.coverage95["pit"]) == (ref.mape["pit"], ref.coverage80["pit"], ref.coverage95["pit"])
    # other taxes keep their 2025 targets
    vat_without = dict(data)
    vat_without["rev_vat"] = data["rev_vat"].loc[:2024]
    assert backtest(vat_without, horizon=2).mape["vat"] != res.mape["vat"]


def test_backtest_does_not_score_or_train_on_2025_pit(monkeypatch):
    import app.model.backtest as bt
    from app.data.dataset import assemble
    from tests.real_dataset import real_results

    data = {k: df.copy() for k, (_, df) in assemble(real_results()).items()}
    for v in bt.MACRO_VARS:  # macro through 2026 so that origins >= 2025 exist
        df = data[v]
        data[v] = df.reindex(range(int(df.index.min()), 2027)).ffill()
    seen = []
    real_sim = bt.simulate

    def spy(model, train, *a, **k):
        res = real_sim(model, train, *a, **k)
        seen.append((int(train["rev_pit"].index.max()), res.state_share.s0, int(train["wage_nominal"].index.max())))  # (rev_pit max, s0, origin)
        return res

    monkeypatch.setattr(bt, "simulate", spy)
    res = bt.backtest(data, horizon=2)
    assert "pit" in res.mape and seen
    assert any(origin >= 2025 for _, _, origin in seen)  # the training window reaches the break year
    assert all(mx < 2025 and s0 == 1.0 for mx, s0, _wage_max in seen)


def test_backtest_note_mentions_share():
    from app.model.backtest import NOTE_PIT_BREAK

    assert "udział budżetu państwa po reformie" in NOTE_PIT_BREAK
    assert "oszacowany z jednego roku (2025)" in NOTE_PIT_BREAK
    assert "nie do oceny backtestu" in NOTE_PIT_BREAK
    from app.model.simulate import METHOD_CHOICES

    assert any(k == "STRUCTURAL_BREAKS" and "udział" in t for k, t in METHOD_CHOICES)


def _share_result(client, mult=0.8):
    c, store = client
    hist, model = fill(store)
    res = sim_result(hist, model)
    res["assumptions"] = [
        *res.get("assumptions", []),
        {"key": "PIT_STATE_SHARE", "value": 0.19812, "status": "calibrated", "year": 2025, "note": "x"},
    ]
    return c, res, {"pit_state_share_mult": mult}


def test_export_contains_pit_share_metadata(client):
    import csv as _csv
    import io

    from openpyxl import load_workbook

    c, res, params = _share_result(client)
    want = "0,198; mnożnik 0,80; efektywny 0,158; 2025"
    r = c.post("/export?format=csv", json={"params": params, "result": res})
    assert r.status_code == 200
    meta = {row[0]: row[1] for row in _csv.reader(io.StringIO(r.content.decode("utf-8-sig"))) if len(row) == 2}
    assert meta["pit_udzial_budzetu"] == want and want[0] not in "=+-@"
    r = c.post("/export?format=xlsx", json={"params": params, "result": res})
    ws = load_workbook(io.BytesIO(r.content))["Metadane"]
    assert {row[0].value: row[1].value for row in ws.iter_rows()}["pit_udzial_budzetu"] == want


def test_export_without_share_row_still_works(client):
    c, store = client
    hist, model = fill(store)
    res = sim_result(hist, model)
    res["assumptions"] = [a for a in res["assumptions"] if a["key"] != "PIT_STATE_SHARE"]
    for fmt in ("csv", "xlsx"):
        r = c.post(f"/export?format={fmt}", json={"params": {}, "result": res})
        assert r.status_code == 200
    assert "pit_udzial_budzetu" not in c.post("/export?format=csv", json={"params": {}, "result": res}).content.decode("utf-8-sig")


@pytest.mark.parametrize("bad", ["x", True, None, [1], 0, -0.2])
def test_export_share_row_value_validated(client, bad):
    c, res, params = _share_result(client)
    res["assumptions"][-1]["value"] = bad
    r = c.post("/export?format=csv", json={"params": params, "result": res})
    assert r.status_code == 422
